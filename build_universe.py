from openpyxl import load_workbook
from pathlib import Path
import json

source = Path('/home/ubuntu/upload/NSE_FO_Options_Stocks_List.xlsx')
out = Path('/home/ubuntu/Sample-Repo/data/fo_symbols.json')
wb = load_workbook(source, read_only=True, data_only=True)
ws = wb['FO Options Stocks']
symbols = []
for _, name in ws.iter_rows(min_row=2, values_only=True):
    if not name:
        continue
    text = str(name).strip()
    if 'LTIM' in text.upper() or 'L&T INFOTECH' in text.upper():
        continue
    symbols.append(text)
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps({'symbols': symbols, 'excluded': ['LTIMindtree']}, indent=2) + '\n')
print(f'wrote {len(symbols)} symbols to {out}')
